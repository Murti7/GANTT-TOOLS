from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook

from gantt.bc3.models import (
    Capitulo,
    Partida,
    Presupuesto,
    ProjectConfig,
    RecursoElemental,
    RecursoMO,
)
from gantt.reporting.presupuesto_exporter import (
    generar_pres0301_mediciones,
    generar_pres0302_mediciones_ciegas,
)
from gantt.reporting.xlsx_integrity import validate_xlsx_integrity


def _presupuesto_pres03() -> Presupuesto:
    partida = Partida(
        codigo="01.01",
        descripcion="Levantamiento inicial",
        descripcion_larga=(
            "Servicio de recopilacion y caracterizacion tecnica.\n\n"
            "Incluye:\n"
            "- Revision documental\n"
            "- Mediciones en m2 y validacion de datos"
        ),
        unidad="m2",
        precio_unitario=1234.56,
        cantidad=0.125,
        lineas_mo=[],
    )
    subpartida = Partida(
        codigo="01.02",
        descripcion="Inspeccion complementaria",
        descripcion_larga="Descripcion tecnica complementaria.",
        unidad="ud",
        precio_unitario=4321.0,
        cantidad=1.5,
        lineas_mo=[],
    )
    subcapitulo = Capitulo(
        codigo="01.1#",
        descripcion="Subcapitulo tecnico",
        importe_total=9876.54,
        partidas=[subpartida],
        subcapitulos=[],
    )
    capitulo = Capitulo(
        codigo="01#",
        descripcion="Informacion tecnica",
        importe_total=9876.54,
        partidas=[partida],
        subcapitulos=[subcapitulo],
    )
    return Presupuesto(
        codigo="OBRA##",
        descripcion="Presupuesto demo",
        importe_total=9876.54,
        capitulos=[capitulo],
        recursos_mo={
            "MO-01": RecursoMO(
                codigo="MO-01",
                descripcion="Oficial especialista",
                precio_hora=88.88,
            )
        },
        recursos_mt={
            "MT-01": RecursoElemental(
                codigo="MT-01",
                descripcion="Material medible",
                unidad="ud",
                precio_unidad=12.34,
            )
        },
        descompuestos_raw={
            "01.01": [("MO-01", 2.0), ("MT-01", 3.0), ("%MT", 0.05)],
        },
        config=ProjectConfig(
            proyecto="Auditoria energetica integral",
            edificio="Viding Fitness Calvia",
            entidad="Cliente demo",
            numero_expediente="BAF-VID-2026-001",
            revision="00",
        ),
    )


def _values(path: Path) -> tuple[list, str]:
    wb = load_workbook(path, data_only=False)
    try:
        ws = wb.active
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        return values, "\n".join(str(value) for value in values)
    finally:
        wb.close()


def _assert_delivery_printing_and_perimeter(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    try:
        ws = wb.active
        assert ws.freeze_panes is None
        assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_A4)
        assert ws.page_setup.orientation == "portrait"
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_margins.left >= 0.7
        assert ws.print_title_rows
        assert ws.print_area

        header_row = int(ws.print_title_rows.split(":")[0].replace("$", ""))
        if ws["A7"].fill.fgColor.rgb.endswith("1B3A5C"):
            assert ws["A7"].font.color.rgb.endswith("FFFFFF")
        min_col = 1
        max_col = ws.max_column
        assert ws.cell(header_row, min_col).border.left.style
        assert ws.cell(header_row, max_col).border.right.style
        for col in range(min_col, max_col + 1):
            assert ws.cell(header_row, col).border.top.style
            assert ws.cell(header_row, col).border.bottom.style
    finally:
        wb.close()


def test_pres0301_mediciones_ciegas_detalladas_layout_integridad_y_contenido(tmp_path: Path):
    path = tmp_path / "PRES.03.01_Mediciones_Ciegas.xlsx"

    generar_pres0301_mediciones(_presupuesto_pres03(), path)

    report = validate_xlsx_integrity(path)
    assert report.passed, report.issues
    _assert_delivery_printing_and_perimeter(path)

    wb = load_workbook(path, data_only=False)
    try:
        assert wb.sheetnames == ["Mediciones Ciegas"]
        ws = wb["Mediciones Ciegas"]
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        text = "\n".join(str(value) for value in values)

        assert "MEDICIONES CIEGAS" in values
        assert "PRES.03.01" in values
        assert "01.01" in values
        assert "Levantamiento inicial" in values
        assert "Servicio de recopilacion" in text
        assert "MO" in values
        assert "MT" in values
        assert "MO-01" in text
        assert "MT-01" in text
        assert 2.0 in values
        assert 3.0 in values
        assert "P. unit." not in text
        assert "Importe" not in text
        assert "%" not in values
        assert "%MT" not in text
        assert "TOTAL CAP" not in text
        assert "PEM" not in text
    finally:
        wb.close()

    with zipfile.ZipFile(path) as archive:
        serialized = b"\n".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
    assert b"1234.56" not in serialized
    assert b"9876.54" not in serialized
    assert b"88.88" not in serialized


def test_pres0302_mediciones_simples_no_escribe_recursos_ni_totales(tmp_path: Path):
    path = tmp_path / "PRES.03.02_Mediciones.xlsx"

    generar_pres0302_mediciones_ciegas(_presupuesto_pres03(), path)

    report = validate_xlsx_integrity(path)
    assert report.passed, report.issues
    _assert_delivery_printing_and_perimeter(path)

    wb = load_workbook(path, data_only=False)
    try:
        assert wb.sheetnames == ["MEDICIONES"]
        ws = wb["MEDICIONES"]
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        text = "\n".join(str(value) for value in values)

        assert "MEDICIONES" in values
        assert "MEDICIONES CIEGAS" not in values
        assert "PRES.03.02" in values
        assert "01.01" in values
        assert "01.02" in values
        assert "m2" in values
        assert 0.125 in values
        assert 1.5 in values
        assert "Servicio de recopilacion" in text

        forbidden_terms = [
            "P. unit.",
            "Importe",
            "Precio unitario",
            "TOTAL CAP",
            "PEM",
            "PEC",
            "GG",
            "BI",
            "IVA",
            "MO",
            "MT",
            "MQ",
            "PA",
            "%MT",
        ]
        for term in forbidden_terms:
            assert not re.search(rf"(?<!\w){re.escape(term)}(?!\w)", text, flags=re.IGNORECASE)
    finally:
        wb.close()


def test_pres0301_y_pres0302_son_documentos_distintos_desde_mismo_presupuesto(tmp_path: Path):
    presupuesto = _presupuesto_pres03()
    ciegas = tmp_path / "PRES.03.01_Mediciones_Ciegas.xlsx"
    mediciones = tmp_path / "PRES.03.02_Mediciones.xlsx"

    generar_pres0301_mediciones(presupuesto, ciegas)
    generar_pres0302_mediciones_ciegas(presupuesto, mediciones)

    valores_ciegas, texto_ciegas = _values(ciegas)
    valores_mediciones, texto_mediciones = _values(mediciones)

    assert ciegas.name != mediciones.name
    assert "PRES.03.01" in valores_ciegas
    assert "PRES.03.02" in valores_mediciones
    assert "01.01" in valores_ciegas
    assert "01.01" in valores_mediciones
    assert "m2" in valores_ciegas
    assert "m2" in valores_mediciones
    assert 0.125 in valores_ciegas
    assert 0.125 in valores_mediciones
    assert texto_ciegas != texto_mediciones
    assert "MO-01" in texto_ciegas
    assert "MT-01" in texto_ciegas
    assert "MO-01" not in texto_mediciones
    assert "MT-01" not in texto_mediciones
    assert "TOTAL CAP" not in texto_ciegas
    assert "TOTAL CAP" not in texto_mediciones
