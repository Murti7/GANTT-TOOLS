from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from gantt.bc3.models import BrandingConfig, ProjectConfig
from gantt.bc3.parser import parse_bc3
from gantt.billing.calculator import calculate_invoice
from gantt.billing.models import Client, InvoiceLine, Issuer
from gantt.billing.services import create_invoice
from gantt.reporting.documents import (
    BillingMetadata,
    DocumentMetadata,
    DocumentStatus,
    DocumentType,
    filename_for_document,
)
from gantt.reporting.invoice_exporter import export_invoice_excel
from gantt.reporting.presentation import build_presentation_context
from gantt.reporting.presupuesto_exporter import generar_pres0201
from gantt.reporting.rendering import LogoRenderPolicy, calculate_logo_size


BC3_MINIMO_IMPORTADO = r"""~C|OBRA##||Importado desde ficheros CSV|100.00||0|
~D|OBRA##
|A#\\1
\|
~C|A#||Capitulo A|100.00||0|
~D|A#
|A01\\2
\|
~C|A01|ud|Partida de ejemplo|50.00||0|
~D|A01
|MO-OFI1\\3
\|
~C|MO-OFI1|h|Oficial 1a|20.00||1|
"""


def test_logo_scaling_preserva_aspect_ratio():
    size = calculate_logo_size(
        1000,
        250,
        LogoRenderPolicy(container_width=180, container_height=60, max_width=150, max_height=44),
    )

    assert size.width <= 150
    assert size.height <= 44
    assert abs((size.width / size.height) - 4.0) < 0.15


def test_budget_header_no_usa_importado_como_titulo_y_configura_printing(tmp_path: Path):
    bc3 = tmp_path / "mini.bc3"
    bc3.write_text(BC3_MINIMO_IMPORTADO, encoding="latin-1")
    presupuesto = parse_bc3(bc3)
    presupuesto = presupuesto.model_copy(
        update={
            "config": ProjectConfig(
                proyecto="Auditoria energetica integral",
                edificio="Viding Fitness Calvia",
                entidad="Cliente demo",
                numero_expediente="BAF-VID-2026-001",
                revision="00",
            ),
            "company": BrandingConfig(nombre="BAFRAS Engineering S.L."),
        }
    )
    output = tmp_path / "PRES.02.01_Cuadro_Precios_1.xlsx"

    generar_pres0201(presupuesto, output)

    wb = load_workbook(output)
    try:
        ws = wb.active
        top_values = [ws.cell(row, 1).value for row in range(1, 10)]
        assert "IMPORTADO DESDE FICHEROS CSV" not in [str(v).upper() for v in top_values if v]
        assert "AUDITORIA ENERGETICA INTEGRAL" in top_values
        assert any("PRES.02.01" in str(value) for value in top_values if value)
        assert ws.print_area
        assert ws.page_setup.orientation == "portrait"
        assert ws.freeze_panes.startswith("A")
        assert int(ws.freeze_panes[1:]) >= 10
        assert ws.oddFooter.center.text == "BAF-VID-2026-001 - Rev. 00"
    finally:
        wb.close()


def test_invoice_exporter_usa_metadata_footer_printing_y_logo_fallback(tmp_path: Path):
    issuer = Issuer(
        legal_name="David Murti",
        tax_id="12345678Z",
        fiscal_address="Calle Demo 1",
        default_withholding_rate=Decimal("0.15"),
        withholding_enabled=True,
        iban="ES0000000000000000000000",
        account_holder="David Murti",
    )
    client = Client(legal_name="Cliente S.A.", tax_id="A12345678")
    invoice = create_invoice(
        issuer,
        client,
        [InvoiceLine(description="Honorarios profesionales", unit_price=1000)],
        project_name="Proyecto demo",
        project_reference="REF-001",
    )
    metadata = DocumentMetadata(
        document_type=DocumentType.INVOICE,
        title="Factura",
        project_name="Proyecto demo",
        reference="REF-001",
        issue_date=date(2026, 8, 21),
        status=DocumentStatus.DRAFT,
        billing=BillingMetadata(invoice_status="draft", due_date=None),
    )

    output = export_invoice_excel(
        invoice,
        calculate_invoice(invoice),
        build_presentation_context(BrandingConfig(nombre="David Murti")),
        tmp_path / "invoice.xlsx",
        metadata=metadata,
    )

    wb = load_workbook(output)
    try:
        ws = wb["Factura"]
        values = [cell.value for row in ws.iter_rows(min_row=1, max_row=12) for cell in row]
        assert "FACTURA" in values
        assert any("Numero: BORRADOR" == value for value in values)
        assert "Proyecto demo" in values
        assert ws.print_area
        assert ws.page_setup.orientation == "portrait"
        assert ws.oddFooter.center.text == "REF-001 - Rev. 00"
    finally:
        wb.close()


def test_filename_invoice_draft_es_estable_desde_metadata():
    metadata = DocumentMetadata(
        document_type=DocumentType.INVOICE,
        reference="REF-001",
        billing=BillingMetadata(invoice_status="draft"),
    )

    assert filename_for_document(metadata) == "DRAFT_REF-001.xlsx"
