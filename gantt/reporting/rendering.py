"""Primitivas de rendering documental para Excel.

Este modulo contiene geometria, logo, cabecera, footer e impresion. No calcula
valores de dominio ni resuelve configuracion desde filesystem.
"""

from dataclasses import dataclass
from datetime import date

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gantt.bc3.models import BrandingConfig, LogoAssetType, ProjectConfig
from gantt.reporting.documents import (
    BillingMetadata,
    DocumentMetadata,
    DocumentPurpose,
    DocumentType,
    document_purpose,
    document_type_info,
)
from gantt.reporting.presentation import PresentationContext
from gantt.reporting.styles import DocumentPalette, LITERALES, build_palette


THIN_GREY = Side(style="thin", color="D9D9D9")
BOX_BORDER = Border(
    left=THIN_GREY,
    right=THIN_GREY,
    top=THIN_GREY,
    bottom=THIN_GREY,
)


@dataclass(frozen=True)
class LogoRenderPolicy:
    """Politica unica de dimensionado de logo."""

    container_width: int = 230
    container_height: int = 62
    max_width: int = 190
    max_height: int = 52
    alignment: str = "center"
    clear_space: int = 8
    preserve_aspect_ratio: bool = True
    minimum_visual_presence: int = 120


@dataclass(frozen=True)
class RenderedLogoSize:
    width: int
    height: int
    original_width: int
    original_height: int


def _font(palette: DocumentPalette, *, size: int, bold: bool = False, color: str | None = None) -> Font:
    return Font(
        name=palette.font_family,
        size=size,
        bold=bold,
        color=(color or palette.color_texto).lstrip("#"),
    )


def calculate_logo_size(
    original_width: int,
    original_height: int,
    policy: LogoRenderPolicy | None = None,
) -> RenderedLogoSize:
    """Calcula el mayor tamano posible preservando aspecto."""
    policy = policy or LogoRenderPolicy()
    available_width = max(1, min(policy.container_width, policy.max_width) - policy.clear_space)
    available_height = max(1, min(policy.container_height, policy.max_height) - policy.clear_space)
    if original_width <= 0 or original_height <= 0:
        return RenderedLogoSize(0, 0, original_width, original_height)
    scale = min(available_width / original_width, available_height / original_height)
    return RenderedLogoSize(
        width=max(1, round(original_width * scale)),
        height=max(1, round(original_height * scale)),
        original_width=original_width,
        original_height=original_height,
    )


def add_logo(ws, company: BrandingConfig | None, policy: LogoRenderPolicy | None = None) -> RenderedLogoSize | None:
    """Inserta logo si existe y devuelve dimensiones renderizadas."""
    if not company or not company.logo_path:
        return None
    try:
        img = XLImage(str(company.logo_path))
        size = calculate_logo_size(img.width, img.height, policy)
        img.width = size.width
        img.height = size.height
        img.anchor = "A1"
        ws.add_image(img)
        return size
    except Exception:
        return None


def _clean_pending(value: str | None, *, lower: bool = False) -> str:
    text = (value or "").strip()
    if not text or text.upper() == "PENDIENTE":
        return "pendiente" if lower else "-"
    return text


def _unique_non_empty(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = (value or "").strip()
        key = text.casefold()
        if text and key not in seen:
            result.append(text)
            seen.add(key)
    return result


def metadata_from_project_config(
    config: ProjectConfig,
    document_type: DocumentType,
    *,
    title: str = "",
) -> DocumentMetadata:
    """Adaptador legacy: crea DocumentMetadata sin consultar fuentes externas."""
    info = document_type_info(document_type)
    return DocumentMetadata(
        document_type=document_type,
        title=title or info.default_title,
        project_code=config.numero_expediente,
        project_name=config.proyecto,
        site_name=config.edificio,
        client_name=config.entidad,
        reference=config.numero_expediente,
        revision=config.revision,
        issue_date=date.today(),
    )


def render_budget_header(
    ws,
    metadata: DocumentMetadata,
    presentation: PresentationContext,
    num_columns: int,
) -> int:
    """Cabecera compacta para documentos Budget DELIVERY."""
    palette = presentation.palette
    company = presentation.company
    logo_size = add_logo(ws, company, LogoRenderPolicy(max_width=185, max_height=48))
    logo_contains_text = _logo_contains_identity_text(company)
    info = document_type_info(metadata.document_type)

    for row, height in [(1, 24), (2, 20), (3, 8), (4, 14), (5, 24), (6, 20), (7, 24)]:
        ws.row_dimensions[row].height = height

    left_end = min(max(2, num_columns // 2), max(1, num_columns - 2))
    if num_columns >= 4:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=left_end)
        ws.merge_cells(start_row=1, start_column=left_end + 1, end_row=1, end_column=num_columns)
        ws.merge_cells(start_row=2, start_column=left_end + 1, end_row=2, end_column=num_columns)
        ws.cell(1, left_end + 1).value = metadata.reference or metadata.project_code
        ws.cell(2, left_end + 1).value = (
            f"Rev. {metadata.revision} - {metadata.issue_date.strftime('%d/%m/%Y')}"
        )
        ws.cell(1, left_end + 1).font = _font(palette, size=11, bold=True)
        ws.cell(2, left_end + 1).font = _font(palette, size=9)
        ws.cell(1, left_end + 1).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(2, left_end + 1).alignment = Alignment(horizontal="right", vertical="center")
    else:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_columns)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_columns)
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_columns)
        if not logo_contains_text:
            ws.cell(1, 1).value = company.nombre if company else ""
        ws.cell(3, 1).value = (
            f"{metadata.reference or metadata.project_code} - "
            f"Rev. {metadata.revision} - {metadata.issue_date.strftime('%d/%m/%Y')}"
        )
        ws.cell(3, 1).font = _font(palette, size=9)

    if not logo_size:
        ws.cell(1, 1).value = company.nombre if company else ""
        ws.cell(1, 1).font = _font(palette, size=14, bold=True, color=palette.color_primario)
    if not logo_contains_text:
        ws.cell(2, 1).value = _company_claim(company)
        ws.cell(2, 1).font = _font(palette, size=8, color=palette.color_texto)

    project_lines = _unique_non_empty([metadata.client_name, metadata.site_name])
    display_client = (project_lines[0] if project_lines else metadata.project_name).upper()
    display_project = metadata.project_name

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=num_columns)
    ws.cell(4, 1).value = "CLIENTE"
    ws.cell(4, 1).font = _font(palette, size=8, bold=True, color=palette.color_primario)

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=num_columns)
    ws.cell(5, 1).value = display_client
    ws.cell(5, 1).font = _font(palette, size=13, bold=True, color=palette.color_primario)

    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=num_columns)
    ws.cell(6, 1).value = display_project if display_project.casefold() != display_client.casefold() else ""
    ws.cell(6, 1).font = _font(palette, size=10)

    if num_columns >= 4:
        ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=num_columns - 1)
        ws.cell(7, num_columns).value = info.document_code
        ws.cell(7, num_columns).font = palette.font_header
        ws.cell(7, num_columns).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(7, 1).value = (metadata.title or info.default_title).upper()
    else:
        ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=num_columns)
        ws.cell(7, 1).value = f"{(metadata.title or info.default_title).upper()} - {info.document_code}"
    ws.cell(7, 1).fill = palette.fill_header
    ws.cell(7, 1).font = palette.font_header
    ws.cell(7, 1).alignment = Alignment(horizontal="left", vertical="center")

    for row in range(1, 8):
        for col in range(1, num_columns + 1):
            ws.cell(row, col).border = BOX_BORDER
            if row == 7:
                ws.cell(row, col).fill = palette.fill_header
            else:
                ws.cell(row, col).fill = palette.fill_white
            ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
    return 8


def render_analysis_header(
    ws,
    metadata: DocumentMetadata,
    presentation: PresentationContext,
    num_columns: int,
) -> int:
    """Cabecera compacta para libros de analisis en pantalla."""
    palette = presentation.palette
    company = presentation.company
    info = document_type_info(metadata.document_type)
    add_logo(ws, company, LogoRenderPolicy(max_width=120, max_height=30))
    for row, height in [(1, 24), (2, 18), (3, 18), (4, 6), (5, 20)]:
        ws.row_dimensions[row].height = height
    title = "PLANIFICACION" if metadata.document_type == DocumentType.PLANNING_ANALYSIS else "BUDGET ANALYSIS"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_columns)
    ws.cell(1, 1).value = f"{company.nombre if company else ''} - {title}".strip(" -")
    ws.cell(1, 1).fill = palette.fill_header
    ws.cell(1, 1).font = palette.font_header
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_columns)
    ws.cell(2, 1).value = metadata.project_name
    ws.cell(2, 1).font = _font(palette, size=11, bold=True, color=palette.color_primario)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_columns)
    parts = []
    if metadata.planning and metadata.planning.scenario:
        parts.append(f"Escenario: {metadata.planning.scenario}")
    parts.extend([metadata.reference or metadata.project_code, metadata.issue_date.strftime("%d/%m/%Y")])
    ws.cell(3, 1).value = " - ".join(p for p in parts if p)
    ws.cell(3, 1).font = _font(palette, size=9)
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=num_columns)
    ws.cell(5, 1).value = f"{info.document_code} - {metadata.title or info.default_title}"
    ws.cell(5, 1).font = _font(palette, size=9, bold=True)
    for row in range(1, 6):
        for col in range(1, num_columns + 1):
            ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
    return 6


def render_planning_header(
    ws,
    metadata: DocumentMetadata,
    presentation: PresentationContext,
    num_columns: int,
) -> int:
    """Cabecera especifica para planificacion ANALYSIS."""
    metadata = metadata.model_copy(update={"title": metadata.title or "Planificacion del proyecto"})
    return render_analysis_header(ws, metadata, presentation, num_columns)


def render_invoice_header(
    ws,
    metadata: DocumentMetadata,
    presentation: PresentationContext,
    num_columns: int,
) -> int:
    """Cabecera especifica para facturas."""
    palette = presentation.palette
    company = presentation.company
    add_logo(ws, company, LogoRenderPolicy(max_width=180, max_height=48))
    for row, height in [(1, 28), (2, 22), (3, 18), (4, 8)]:
        ws.row_dimensions[row].height = height
    left_end = 3 if num_columns >= 6 else max(1, num_columns // 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=left_end)
    if not company or not company.logo_path:
        ws.cell(1, 1).value = company.nombre if company else ""
        ws.cell(1, 1).font = _font(palette, size=15, bold=True, color=palette.color_primario)

    billing = metadata.billing or BillingMetadata()
    right_start = left_end + 1
    status_or_number = billing.invoice_number or "BORRADOR"
    ws.merge_cells(start_row=1, start_column=right_start, end_row=1, end_column=num_columns)
    ws.merge_cells(start_row=2, start_column=right_start, end_row=2, end_column=num_columns)
    ws.merge_cells(start_row=3, start_column=right_start, end_row=3, end_column=num_columns)
    ws.cell(1, right_start).value = (metadata.title or "Factura").upper()
    ws.cell(2, right_start).value = status_or_number
    ws.cell(3, right_start).value = metadata.issue_date.strftime("%d/%m/%Y")
    if billing.due_date and billing.invoice_status != "draft":
        ws.cell(3, right_start).value = f"{metadata.issue_date.strftime('%d/%m/%Y')} - Vence {billing.due_date.strftime('%d/%m/%Y')}"
    ws.cell(1, right_start).font = _font(palette, size=18, bold=True)
    ws.cell(2, right_start).font = _font(palette, size=12, bold=True, color=palette.color_primario)
    ws.cell(3, right_start).font = _font(palette, size=9)
    for row in (1, 2, 3):
        ws.cell(row, right_start).alignment = Alignment(horizontal="right", vertical="center")
    for row in range(1, 4):
        for col in range(1, num_columns + 1):
            ws.cell(row, col).border = BOX_BORDER
            ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
    return 5


def apply_document_footer(
    ws,
    metadata: DocumentMetadata,
    presentation: PresentationContext,
    num_columns: int,
    *,
    invoice: bool = False,
) -> None:
    """Footer visible y footer de impresion."""
    palette = presentation.palette
    company = presentation.company
    col_max = get_column_letter(num_columns)
    row = ws.max_row + 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_columns)
    if company:
        if invoice:
            text = " · ".join(
                p for p in [
                    company.nombre,
                    "" if company.nif_cif.upper() == "PENDIENTE" else company.nif_cif,
                    "" if company.direccion_fiscal.upper() == "PENDIENTE" else company.direccion_fiscal,
                    "" if company.web.upper() == "PENDIENTE" else company.web,
                ] if p
            )
        else:
            text = " | ".join(
                p for p in [
                    f"{company.nombre} - {_clean_pending(company.web)}" if company.web else company.nombre,
                    f"{metadata.reference or metadata.project_code} - Rev. {metadata.revision}",
                ] if p
            )
    else:
        text = f"{metadata.reference or metadata.project_code} - Rev. {metadata.revision}"
    c = ws.cell(row, 1)
    c.value = text
    c.font = palette.font_small
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = palette.fill_alt
    ws.row_dimensions[row].height = 16
    ws.oddFooter.left.text = (company.nombre if company else "")
    ws.oddFooter.center.text = f"{metadata.reference or metadata.project_code} - Rev. {metadata.revision}"
    ws.oddFooter.right.text = "Pagina &[Page] de &[Pages]"
    ws.print_area = f"$A$1:${col_max}${ws.max_row}"


def configure_excel_printing(
    ws,
    *,
    orientation: str = "portrait",
    repeat_row: int | None = None,
    fit_to_height: int = 0,
    purpose: DocumentPurpose = DocumentPurpose.DELIVERY,
) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_to_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.25
    ws.page_margins.footer = 0.25
    ws.sheet_view.showGridLines = False
    ws.sheet_view.view = "pageLayout"
    if repeat_row is not None:
        ws.print_title_rows = f"{repeat_row}:{repeat_row}"


def _company_claim(company: BrandingConfig | None) -> str:
    if not company:
        return ""
    parts = [company.web, company.email, company.telefono]
    contact = " | ".join(p for p in parts if p and p != "PENDIENTE")
    return contact or "Ingenieria - Energia - Sistemas"


def _logo_contains_identity_text(company: BrandingConfig | None) -> bool:
    if not company or not company.logo_path:
        return False
    return company.logo_asset_type in {
        LogoAssetType.WORDMARK,
        LogoAssetType.LOCKUP_HORIZONTAL,
        LogoAssetType.LOCKUP_VERTICAL,
    }


def presentation_from_palette_company(
    palette: DocumentPalette | None,
    company: BrandingConfig | None,
) -> PresentationContext:
    pal = palette or build_palette(company)
    return PresentationContext(
        company=company,
        palette=pal,
        language=company.idioma if company else "es",
        theme_id=(company.nombre.lower().replace(" ", "-") if company else "default"),
    )
