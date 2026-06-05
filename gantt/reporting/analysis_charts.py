"""
Módulo responsable de generar gráficos estáticos de análisis del presupuesto.

Responsabilidad: leer el Excel de análisis ya generado y producir representaciones
visuales PNG organizadas en subcarpetas dentro del directorio de output del proyecto.
No depende del modelo Presupuesto ni del BC3. Opera únicamente sobre el Excel.
"""

import re
import textwrap
import unicodedata
from pathlib import Path

import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.cm as cm
from openpyxl import load_workbook

from gantt.reporting.styles import CHART_DPI, CHART_PALETA, CHART_PALETA_APILADA
from gantt.reporting.palette import (
    CHART_AZUL, CHART_AZUL_MEDIO, CHART_VERDE,
    CHART_GRIS_CLARO, CHART_BORDE, CHART_GRID,
    CHART_TEXTO, CHART_MUTED, FONT_CHART,
)

# Alias locales para compatibilidad con el código existente del módulo
COLOR_BLUE       = CHART_AZUL
from gantt.reporting.palette import CHART_AZUL_CLARO
COLOR_BLUE_LIGHT = CHART_AZUL_CLARO
COLOR_GREEN      = CHART_VERDE
COLOR_GRID       = CHART_GRID
COLOR_TEXT       = CHART_TEXTO
COLOR_MUTED      = CHART_MUTED

plt.rcParams['font.family']     = FONT_CHART
plt.rcParams['axes.edgecolor']  = CHART_TEXTO
plt.rcParams['axes.labelcolor'] = CHART_TEXTO
plt.rcParams['xtick.color']     = CHART_TEXTO
plt.rcParams['ytick.color']     = CHART_TEXTO


# ---------------------------------------------------------------------------
# Funciones auxiliares
# ---------------------------------------------------------------------------

def clean_label(text: str) -> str:
    """
    Limpia una cadena para visualización en matplotlib.
    """
    replacements = {
        '': '-',
        '–': '-',
        '—': '-',
        '’': "'",
        'ª': 'a',
        'º': 'o',
    }
    result = str(text)
    for char, replacement in replacements.items():
        result = result.replace(char, replacement)
    return result


def normalize_col(name: str) -> str:
    """
    Normaliza un nombre de columna para comparaciones robustas.
    """
    if not name:
        return ''
    nfkd = unicodedata.normalize('NFKD', str(name))
    ascii_str = nfkd.encode('ascii', 'ignore').decode('ascii')
    cleaned = re.sub(r'[^a-z0-9 ]', ' ', ascii_str.lower())
    return re.sub(r'\s+', ' ', cleaned).strip()


def find_column(columns: list, aliases: list[str]) -> str | None:
    """
    Busca una columna por alias normalizados.
    """
    norm_aliases = {normalize_col(a) for a in aliases}
    for col in columns:
        if col is not None and normalize_col(str(col)) in norm_aliases:
            return col
    return None


def format_eur(value: float) -> str:
    """
    Formato europeo sin decimales para importes.
    """
    return f'{value:,.0f}'.replace(',', '.') + ' €'


def wrap_label(text: str, width: int = 42, max_lines: int = 2) -> str:
    """
    Parte una etiqueta larga en un máximo de líneas.
    """
    text = clean_label(text)
    lines = textwrap.wrap(text, width=width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip('.') + '...'
    return '\n'.join(lines)


def simplify_profile_name(name: str) -> str:
    """
    Simplifica nombres largos de perfiles para mejorar legibilidad gráfica.
    """
    replacements = {
        'Técnico especialista control': 'Técnico control senior',
        'Tecnico especialista control': 'Técnico control senior',
        'Técnico junior control': 'Técnico control junior',
        'Tecnico junior control': 'Técnico control junior',
        'Oficial 1ª climatización': 'Oficial climatización',
        'Oficial 1a climatización': 'Oficial climatización',
        'Ayudante climatización': 'Ayudante clima',
        'Oficial 1ª fontanero': 'Oficial fontanería',
        'Oficial 1a fontanero': 'Oficial fontanería',
        'Ayudante de fontanero': 'Ayudante fontanería',
        'Oficial 1ª electricista': 'Oficial electricista',
        'Oficial 1a electricista': 'Oficial electricista',
        'Ayudante electricista': 'Ayudante electricista',
        'Oficial 1ª construcción de obra civil.': 'Oficial obra civil',
        'Oficial 1a construcción de obra civil.': 'Oficial obra civil',
        'Ayudante construcción de obra civil.': 'Ayudante obra civil',
        'Delineante de 1º': 'Delineante',
        'Delineante de 1o': 'Delineante',
    }
    return replacements.get(str(name), str(name))


def setup_axes(ax) -> None:
    """
    Aplica estilo base común a los ejes.
    """
    ax.set_facecolor('white')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['bottom'].set_visible(False)
    ax.grid(axis='x', linestyle='--', alpha=0.16, color=COLOR_GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def save_barh_chart(
    values: list[float],
    labels: list[str],
    path: Path,
    title: str,
    xlabel: str,
    color: str = COLOR_BLUE,
    value_fmt: str = '{:,.1f}',
    wrap_labels: bool = True,
    filter_zero: bool = True,
) -> None:
    """
    Genera y guarda un gráfico de barras horizontales.
    """
    rows = [
        (label, float(value or 0))
        for label, value in zip(labels, values)
        if not filter_zero or float(value or 0) > 0
    ]

    if not rows:
        return

    labels = [r[0] for r in rows]
    values = [r[1] for r in rows]

    n = len(labels)
    max_val = max(values) if values else 1
    fig_height = max(4.2, n * 0.52 + 1.2)

    fig, ax = plt.subplots(figsize=(12.5, fig_height))
    fig.patch.set_facecolor('white')

    y_pos = list(range(n))
    bars = ax.barh(y_pos, values, color=color, edgecolor='white', height=0.54)

    rendered_labels = [
        wrap_label(label, width=48, max_lines=2) if wrap_labels else clean_label(label)
        for label in labels
    ]

    ax.set_yticks(y_pos)
    ax.set_yticklabels(rendered_labels, fontsize=9)
    ax.invert_yaxis()
    ax.set_xlabel(xlabel, fontsize=10)
    ax.set_title(title, fontsize=16, fontweight='bold', pad=16)
    ax.set_xlim(0, max_val * 1.10)

    setup_axes(ax)

    for bar, val in zip(bars, values):
        if val > 0:
            ax.text(
                bar.get_width() + max_val * 0.006,
                bar.get_y() + bar.get_height() / 2,
                value_fmt.format(val),
                va='center',
                ha='left',
                fontsize=8.5,
                color=COLOR_TEXT,
            )

    fig.tight_layout(pad=1.8)
    fig.savefig(path, dpi=CHART_DPI, bbox_inches='tight', facecolor='white')
    plt.close(fig)


def save_top_partidas_chart(
    rows: list[dict],
    col_cod: str | None,
    col_desc: str,
    col_imp: str,
    path: Path,
) -> None:
    """
    Genera el gráfico de top partidas por impacto económico.
    """
    rows = [r for r in rows if float(r.get(col_imp) or 0) > 0]
    if not rows:
        return

    n = len(rows)
    values = [float(r.get(col_imp) or 0) for r in rows]
    max_val = max(values) if values else 1

    labels = []
    for r in rows:
        code = clean_label(str(r.get(col_cod) or '')) if col_cod else ''
        desc = wrap_label(str(r.get(col_desc) or ''), width=46, max_lines=2)
        labels.append((code, desc))

    fig_height = max(6.0, n * 0.58 + 1.7)
    fig, ax = plt.subplots(figsize=(14.5, fig_height))
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    y_pos = list(range(n))
    bars = ax.barh(y_pos, values, color=COLOR_BLUE, edgecolor='white', height=0.52)

    ax.set_yticks(y_pos)
    ax.set_yticklabels(['' for _ in y_pos])
    ax.invert_yaxis()
    ax.set_xlabel('Importe (€ PEM)', fontsize=10)
    ax.set_title('Partidas con mayor impacto económico (€ PEM)', fontsize=16, fontweight='bold', pad=16)
    ax.set_xlim(0, max_val * 1.13)

    setup_axes(ax)

    # Columnas visuales de código y descripción.
    x_code = -max_val * 0.67
    x_desc = -max_val * 0.55

    for i, (code, desc) in enumerate(labels):
        ax.text(
            x_code,
            i,
            code,
            va='center',
            ha='left',
            fontsize=8.7,
            color=COLOR_MUTED,
            fontweight='bold',
            clip_on=False,
        )
        ax.text(
            x_desc,
            i,
            desc,
            va='center',
            ha='left',
            fontsize=8.7,
            color=COLOR_TEXT,
            clip_on=False,
        )

    for bar, val in zip(bars, values):
        ax.text(
            bar.get_width() + max_val * 0.006,
            bar.get_y() + bar.get_height() / 2,
            format_eur(val),
            va='center',
            ha='left',
            fontsize=8.5,
            color=COLOR_TEXT,
        )

    fig.subplots_adjust(left=0.38, right=0.96, top=0.91, bottom=0.09)
    fig.savefig(path, dpi=CHART_DPI, bbox_inches='tight', facecolor='white')
    plt.close(fig)


def save_stacked_mo_chart(
    caps: list[dict],
    profiles: list[str],
    col_desc: str,
    path: Path,
) -> None:
    """
    Genera un gráfico apilado de horas MO por perfil y capítulo.

    Mantiene compatibilidad con el gráfico original, pero:
    - filtra perfiles sin horas;
    - agrupa perfiles menores en 'Otros';
    - usa paleta más controlada;
    - reduce saturación visual.
    """
    caps = [
        r for r in caps
        if sum(float(r.get(col) or 0) for col in profiles) > 0
    ]
    if not caps:
        return

    profile_totals = {
        col: sum(float(r.get(col) or 0) for r in caps)
        for col in profiles
    }

    main_profiles = [
        col for col, total in sorted(profile_totals.items(), key=lambda x: x[1], reverse=True)
        if total >= 50
    ]
    small_profiles = [
        col for col, total in profile_totals.items()
        if total > 0 and total < 50
    ]

    plot_profiles = main_profiles.copy()
    include_otros = bool(small_profiles)

    n_caps = len(caps)
    y_pos = list(range(n_caps))
    left = [0.0] * n_caps

    palette = CHART_PALETA_APILADA

    fig, ax = plt.subplots(figsize=(13.5, max(4.5, n_caps * 0.58 + 1.5)))
    fig.patch.set_facecolor('white')

    for i, col in enumerate(plot_profiles):
        vals = [float(r.get(col) or 0) for r in caps]
        ax.barh(
            y_pos,
            vals,
            left=left,
            label=simplify_profile_name(col),
            color=palette[i % len(palette)],
            height=0.58,
            edgecolor='white',
        )
        left = [l + v for l, v in zip(left, vals)]

    if include_otros:
        vals = [
            sum(float(r.get(col) or 0) for col in small_profiles)
            for r in caps
        ]
        ax.barh(
            y_pos,
            vals,
            left=left,
            label='Otros perfiles (<50 h)',
            color=CHART_GRIS_CLARO,
            height=0.58,
            edgecolor='white',
        )

    ax.set_yticks(y_pos)
    ax.set_yticklabels(
        [wrap_label(str(r.get(col_desc, '')), width=42, max_lines=2) for r in caps],
        fontsize=9,
    )
    ax.invert_yaxis()
    ax.set_xlabel('Horas MO', fontsize=10)
    ax.set_title('Distribución de horas de MO por capítulo y perfil', fontsize=16, fontweight='bold', pad=16)

    setup_axes(ax)

    ax.legend(
        loc='lower right',
        fontsize=8,
        ncol=2,
        framealpha=0.95,
        facecolor='white',
        edgecolor=CHART_BORDE,
    )

    fig.tight_layout(pad=1.8)
    fig.savefig(path, dpi=CHART_DPI, bbox_inches='tight', facecolor='white')
    plt.close(fig)


_RESUMEN_KNOWN_NORMS = frozenset(normalize_col(c) for c in [
    'Código', 'Codigo',
    'Descripción', 'Descripcion',
    'Importe total', 'Importe total (€)',
    'H MO total', 'H MO',
    '% MO / Importe', '% MO',
])


# ---------------------------------------------------------------------------
# Clase principal
# ---------------------------------------------------------------------------

class AnalysisChartsReport:
    """
    Genera gráficos estáticos de análisis económico y de mano de obra
    a partir del Excel de análisis del presupuesto BC3.
    """

    def __init__(self, excel_path: Path, output_dir: Path) -> None:
        self.excel_path = excel_path
        self.output_dir = output_dir
        self.resumen_headers: list = []
        self.resumen_rows: list[dict] = []
        self.mo_profile_cols: list[str] = []
        self.partidas_headers: list = []
        self.partidas_rows: list[dict] = []

    def generate(self) -> None:
        """
        Crea los directorios de salida, lee el Excel y genera todos los gráficos.
        """
        (self.output_dir / '01_economico').mkdir(parents=True, exist_ok=True)
        (self.output_dir / '02_mano_obra').mkdir(parents=True, exist_ok=True)

        self.read_workbook_tables()
        self.write_economic_charts()
        self.write_labor_charts()

    def read_workbook_tables(self) -> None:
        """
        Lee las hojas necesarias del Excel.
        """
        resumen_sheet = 'Resumen por capítulo'
        partidas_sheet = 'Partidas detalladas'

        wb = load_workbook(self.excel_path, read_only=True, data_only=True)

        if resumen_sheet not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Hoja '{resumen_sheet}' no encontrada en {self.excel_path}")

        ws = wb[resumen_sheet]
        rows = list(ws.rows)
        self.resumen_headers = [cell.value for cell in rows[0]] if rows else []

        col_codigo = find_column(self.resumen_headers, ['Código', 'Codigo'])
        if col_codigo is None:
            wb.close()
            raise ValueError("No se encontró columna 'Código' en la hoja Resumen por capítulo")

        idx_codigo = self.resumen_headers.index(col_codigo)

        self.resumen_rows = [
            {
                self.resumen_headers[i]: cell.value
                for i, cell in enumerate(row)
                if i < len(self.resumen_headers)
            }
            for row in rows[1:]
            if row[idx_codigo].value is not None
            and str(row[idx_codigo].value).endswith('#')
        ]

        self.mo_profile_cols = [
            h for h in self.resumen_headers
            if h is not None and normalize_col(str(h)) not in _RESUMEN_KNOWN_NORMS
        ]

        if partidas_sheet in wb.sheetnames:
            ws_p = wb[partidas_sheet]
            rows_p = list(ws_p.rows)
            self.partidas_headers = [cell.value for cell in rows_p[0]] if rows_p else []
            self.partidas_rows = [
                {
                    self.partidas_headers[i]: cell.value
                    for i, cell in enumerate(row)
                    if i < len(self.partidas_headers)
                }
                for row in rows_p[1:]
                if any(cell.value is not None for cell in row)
            ]
        else:
            print(f"      Aviso: hoja '{partidas_sheet}' no encontrada")

        wb.close()

    def get_main_caps(self) -> list[dict]:
        """
        Devuelve solo capítulos principales.

        Importante: no se deben sumar capítulos principales + subcapítulos,
        porque el Excel ya acumula recursivamente los importes y horas en
        los capítulos principales.
        """
        col_cod = find_column(self.resumen_headers, ['Código', 'Codigo'])
        if col_cod is None:
            return []

        return [
            r for r in self.resumen_rows
            if r.get(col_cod) and '.' not in str(r[col_cod])
        ]

    def write_economic_charts(self) -> None:
        """
        Genera gráficos económicos.
        """
        ec_dir = self.output_dir / '01_economico'

        col_cod = find_column(self.resumen_headers, ['Código', 'Codigo'])
        col_desc = find_column(self.resumen_headers, ['Descripción', 'Descripcion'])
        col_imp = find_column(self.resumen_headers, ['Importe total (€)', 'Importe total'])

        main_caps = self.get_main_caps()

        missing = [
            name for name, col in [
                ('Código', col_cod),
                ('Descripción', col_desc),
                ('Importe total (€)', col_imp),
            ]
            if col is None
        ]

        if missing:
            print(f'      Omitidos gráficos económicos por capítulo: faltan columnas {missing}')
        else:
            caps_sorted = sorted(main_caps, key=lambda r: float(r[col_imp] or 0), reverse=True)

            save_barh_chart(
                values=[float(r[col_imp] or 0) for r in caps_sorted],
                labels=[str(r[col_desc]) for r in caps_sorted],
                path=ec_dir / '01_importe_por_capitulo.png',
                title='Importe por capítulo',
                xlabel='Importe (€)',
                value_fmt='{:,.0f}',
                color=COLOR_BLUE,
            )
            print('      OK 01_importe_por_capitulo.png')

            total_imp = sum(float(r[col_imp] or 0) for r in main_caps)
            if total_imp > 0:
                save_barh_chart(
                    values=[float(r[col_imp] or 0) / total_imp * 100 for r in caps_sorted],
                    labels=[str(r[col_desc]) for r in caps_sorted],
                    path=ec_dir / '02_peso_economico_por_capitulo.png',
                    title='Peso económico por capítulo (% sobre PEM)',
                    xlabel='% sobre importe total',
                    value_fmt='{:.1f}%',
                    color=COLOR_GREEN,
                )
                print('      OK 02_peso_economico_por_capitulo.png')
            else:
                print('      Omitido 02_peso_economico_por_capitulo: importe total = 0')

        if not self.partidas_rows:
            print('      Omitido 03_top_partidas_por_importe: hoja Partidas detalladas no disponible')
            return

        col_pcod = find_column(self.partidas_headers, ['Código partida', 'Código', 'Codigo'])
        col_pdesc = find_column(self.partidas_headers, ['Descripción', 'Descripcion', 'Resumen'])
        col_pimp = find_column(self.partidas_headers, ['Importe (€)', 'Importe total (€)', 'Importe'])

        missing_p = [
            name for name, col in [
                ('Descripción', col_pdesc),
                ('Importe (€)', col_pimp),
            ]
            if col is None
        ]

        if missing_p:
            print(f'      Omitido 03_top_partidas_por_importe: faltan columnas {missing_p}')
            return

        top15 = sorted(
            self.partidas_rows,
            key=lambda r: float(r.get(col_pimp) or 0),
            reverse=True,
        )[:15]

        save_top_partidas_chart(
            rows=top15,
            col_cod=col_pcod,
            col_desc=col_pdesc,
            col_imp=col_pimp,
            path=ec_dir / '03_top_partidas_por_importe.png',
        )
        print('      OK 03_top_partidas_por_importe.png')

    def write_labor_charts(self) -> None:
        """
        Genera gráficos de mano de obra.
        """
        mo_dir = self.output_dir / '02_mano_obra'

        col_cod = find_column(self.resumen_headers, ['Código', 'Codigo'])
        col_desc = find_column(self.resumen_headers, ['Descripción', 'Descripcion'])
        col_hmo = find_column(self.resumen_headers, ['H MO total', 'H MO'])

        main_caps = self.get_main_caps()

        missing = [
            name for name, col in [
                ('Código', col_cod),
                ('Descripción', col_desc),
                ('H MO total', col_hmo),
            ]
            if col is None
        ]

        if missing:
            print(f'      Omitido 01_horas_mo_por_capitulo: faltan columnas {missing}')
        else:
            caps_sorted = sorted(main_caps, key=lambda r: float(r.get(col_hmo) or 0), reverse=True)

            save_barh_chart(
                values=[float(r.get(col_hmo) or 0) for r in caps_sorted],
                labels=[str(r[col_desc]) for r in caps_sorted],
                path=mo_dir / '01_horas_mo_por_capitulo.png',
                title='Horas de mano de obra por capítulo',
                xlabel='Horas MO',
                value_fmt='{:,.1f}',
                color=CHART_AZUL_MEDIO,
                filter_zero=True,
            )
            print('      OK 01_horas_mo_por_capitulo.png')

        if not self.mo_profile_cols:
            print('      Omitido 02_horas_por_perfil: no se detectaron columnas de perfiles MO')
            return

        # Importante: sumar solo capítulos principales para no duplicar horas.
        profile_totals = {
            col: sum(float(r.get(col) or 0) for r in main_caps)
            for col in self.mo_profile_cols
        }

        active = {k: v for k, v in profile_totals.items() if v > 0}
        if not active:
            print('      Omitido 02_horas_por_perfil: todos los perfiles tienen 0 horas')
        else:
            sorted_profiles = sorted(active.items(), key=lambda x: x[1], reverse=True)

            save_barh_chart(
                values=[v for _, v in sorted_profiles],
                labels=[simplify_profile_name(k) for k, _ in sorted_profiles],
                path=mo_dir / '02_horas_por_perfil.png',
                title='Total de horas por perfil de mano de obra',
                xlabel='Horas',
                value_fmt='{:,.1f}',
                color=COLOR_BLUE_LIGHT,
                filter_zero=True,
            )
            print('      OK 02_horas_por_perfil.png')

            if col_hmo:
                total_hmo_caps = sum(float(r.get(col_hmo) or 0) for r in main_caps)
                total_profiles = sum(profile_totals.values())
                print(f'      Control horas MO capítulos: {total_hmo_caps:.1f} h')
                print(f'      Control horas MO perfiles:  {total_profiles:.1f} h')

        if not self.mo_profile_cols or not main_caps or col_desc is None:
            missing_parts = (
                (['columnas de perfiles MO'] if not self.mo_profile_cols else [])
                + (['capítulos principales'] if not main_caps else [])
                + (['Descripción'] if col_desc is None else [])
            )
            print(f'      Omitido 03_horas_por_perfil_y_capitulo: faltan {missing_parts}')
            return

        caps_for_stack = sorted(
            main_caps,
            key=lambda r: float(r.get(col_hmo) or 0) if col_hmo else 0,
            reverse=True,
        )

        active_profiles = [
            col for col in self.mo_profile_cols
            if any(float(r.get(col) or 0) > 0 for r in caps_for_stack)
        ]

        if not active_profiles:
            print('      Omitido 03_horas_por_perfil_y_capitulo: sin horas en capítulos principales')
            return

        save_stacked_mo_chart(
            caps=caps_for_stack,
            profiles=active_profiles,
            col_desc=col_desc,
            path=mo_dir / '03_horas_por_perfil_y_capitulo.png',
        )
        print('      OK 03_horas_por_perfil_y_capitulo.png')