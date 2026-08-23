from decimal import Decimal
import json
from pathlib import Path

from openpyxl import load_workbook

from gantt.application.billing import export_invoice_artifacts
from gantt.billing.calculator import calculate_invoice
from gantt.billing.models import Client, ClientType, InvoiceLine, Issuer
from gantt.billing.persistence import write_invoice_json
from gantt.billing.services import create_invoice
from gantt.reporting.invoice_exporter import export_invoice_excel
from gantt.reporting.presentation import build_presentation_context


def test_invoice_exporter_escribe_factura_excel_y_json(tmp_path: Path):
    issuer = Issuer(
        legal_name="David Murti",
        entity_type="professional",
        tax_id="12345678Z",
        fiscal_address="Calle Demo 1",
        default_withholding_rate=Decimal("0.15"),
        withholding_enabled=True,
        iban="ES0000000000000000000000",
        account_holder="David Murti",
    )
    client = Client(
        legal_name="Cliente S.A.",
        client_type=ClientType.COMPANY,
        tax_id="A12345678",
        fiscal_address="Avenida Cliente 2",
    )
    invoice = create_invoice(
        issuer,
        client,
        [InvoiceLine(description="Honorarios profesionales", unit_price=1000)],
        project_name="Proyecto demo",
    )
    calc = calculate_invoice(invoice)
    context = build_presentation_context()

    excel_path = export_invoice_excel(invoice, calc, context, tmp_path / "factura.xlsx")
    json_path = write_invoice_json(invoice, tmp_path / "invoice.json")

    assert excel_path.exists()
    assert json_path.exists()

    wb = load_workbook(excel_path, data_only=True)
    try:
        ws = wb["Factura"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        assert "David Murti" in values
        assert "Cliente S.A." in values
        assert "TOTAL FACTURA" in values
        assert "A PAGAR" in values
        assert values.count("BORRADOR") == 1
        assert "DRAFT" not in values
        assert ws.freeze_panes is None
        assert Decimal("1210.00") in values
        assert Decimal("1060.00") in values
        assert "ES0000000000000000000000" in values
    finally:
        wb.close()


def test_export_invoice_artifacts_escribe_manifest(tmp_path: Path):
    issuer = Issuer(
        legal_name="BAFRAS Engineering S.L.",
        entity_type="company",
        tax_id="B12345678",
        fiscal_address="Calle Demo 1",
    )
    client = Client(legal_name="Cliente S.A.", client_type=ClientType.COMPANY, tax_id="A12345678")
    invoice = create_invoice(
        issuer,
        client,
        [InvoiceLine(description="Factura proyecto", unit_price=500)],
        project_name="Viding Fitness",
    )

    result = export_invoice_artifacts(
        invoice,
        build_presentation_context(),
        tmp_path,
        basename="draft_invoice",
        source="manual-test",
    )

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))

    assert result.excel_path.exists()
    assert result.json_path.exists()
    assert manifest["invoice_id"] == invoice.invoice_id
    assert manifest["invoice_status"] == "draft"
    assert manifest["issuer"] == "BAFRAS Engineering S.L."
    assert manifest["client"] == "Cliente S.A."
    assert manifest["source"] == "manual-test"
    assert manifest["ready_to_issue"] is False
    assert manifest["completeness_status"] == "draft_valid"
    assert "client.address" in manifest["missing_required_fields"]
    assert len(manifest["data_hash"]) == 64
