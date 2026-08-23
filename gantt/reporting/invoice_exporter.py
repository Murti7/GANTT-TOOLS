"""Exporter Excel para facturas.

El exporter representa un Invoice ya calculado. No contiene logica fiscal.
"""

from pathlib import Path
from copy import copy

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gantt.billing.models import Invoice, InvoiceCalculation, TaxTreatment
from gantt.billing.presets import billing_preset
from gantt.reporting.documents import BillingMetadata, DocumentMetadata, DocumentStatus, DocumentType
from gantt.reporting.presentation import PresentationContext
from gantt.reporting.rendering import (
    apply_document_footer,
    configure_excel_printing,
    render_invoice_header,
)


THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
MONEY_FMT = '#.##0,00 [$€-es-ES]'
QTY_FMT = "#,##0.###"
PCT_FMT = "0.00%"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))


def _write_label_value(ws, row: int, label: str, value, value_col: int = 2) -> None:
    ws.cell(row, 1).value = label
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, value_col).value = value


def _section(ws, row: int, title: str, end_col: int, context: PresentationContext) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    cell = ws.cell(row, 1)
    cell.value = title
    cell.fill = context.palette.fill_header
    cell.font = context.palette.font_header
    cell.alignment = Alignment(horizontal="center")
    return row + 1


def _format_money_cell(cell) -> None:
    cell.number_format = MONEY_FMT
    cell.alignment = Alignment(horizontal="right")


def _visible(value) -> str:
    text = "" if value is None else str(value).strip()
    return "-" if not text or text.upper() == "PENDIENTE" else text


def _write_pair_block(ws, row: int, start_col: int, title: str, rows: list[tuple[str, object]], context: PresentationContext) -> None:
    end_col = start_col + 2
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    title_cell = ws.cell(row, start_col)
    title_cell.value = title
    title_cell.fill = context.palette.fill_header
    title_cell.font = context.palette.font_header
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
    visible_rows = [(label, value) for label, value in rows if label == "Nombre" or _visible(value) != "-"]
    for offset, (label, value) in enumerate(visible_rows, start=1):
        label_cell = ws.cell(row + offset, start_col)
        value_cell = ws.cell(row + offset, start_col + 1)
        label_cell.value = label
        label_cell.font = context.palette.font_small
        value_cell.value = _visible(value)
        ws.merge_cells(start_row=row + offset, start_column=start_col + 1, end_row=row + offset, end_column=end_col)
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def export_invoice_excel(
    invoice: Invoice,
    calculation: InvoiceCalculation,
    context: PresentationContext,
    output_path: Path,
    metadata: DocumentMetadata | None = None,
) -> Path:
    """Genera el Excel profesional de factura desde dominio + calculo."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Factura"

    for col, width in {
        "A": 18,
        "B": 42,
        "C": 12,
        "D": 14,
        "E": 16,
        "F": 16,
    }.items():
        ws.column_dimensions[col].width = width

    metadata = metadata or DocumentMetadata(
        document_type=DocumentType.INVOICE,
        title=billing_preset(invoice.billing_preset).document_title,
        project_name=invoice.project_name,
        reference=invoice.project_reference,
        status=DocumentStatus.ISSUED if invoice.status.value == "issued" else DocumentStatus.DRAFT,
        billing=BillingMetadata(
            invoice_number=invoice.invoice_number,
            invoice_status=invoice.status.value,
            due_date=invoice.payment_terms.due_date,
            billing_preset=invoice.billing_preset,
            billing_source_type=invoice.billing_source_type,
            billing_source_reference=invoice.billing_source_reference,
            economic_basis=invoice.economic_basis,
        ),
    )
    row = render_invoice_header(ws, metadata, context, 6)
    issuer_rows = [
        ("Nombre", invoice.issuer.legal_name),
        ("NIF/CIF", invoice.issuer.tax_id),
        ("Direccion", invoice.issuer.fiscal_address),
        ("Email", invoice.issuer.email),
        ("Telefono", invoice.issuer.phone),
    ]
    client_rows = [
        ("Nombre", invoice.client.legal_name),
        ("NIF/CIF", invoice.client.tax_id or invoice.client.vat_id),
        ("Direccion", invoice.client.fiscal_address),
        ("Pais", invoice.client.country),
        ("Email", invoice.client.billing_email),
        ("Pedido", invoice.client.purchase_order or invoice.client.expediente),
    ]
    _write_pair_block(ws, row, 1, "EMISOR", issuer_rows, context)
    _write_pair_block(ws, row, 4, "CLIENTE", client_rows, context)
    row += max(len(issuer_rows), len(client_rows)) + 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row, 1).value = "PROYECTO"
    ws.cell(row, 1).font = context.palette.font_subhead
    ws.cell(row, 1).fill = context.palette.fill_subhead
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    project_parts = [invoice.project_name, metadata.site_name]
    ws.cell(row, 1).value = " - ".join(_visible(p) for p in project_parts if _visible(p) != "-")
    ws.cell(row, 1).font = Font(name=context.palette.font_family, bold=True, size=11)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row, 1).value = f"Ref. {_visible(invoice.project_reference or metadata.reference)}"
    ws.cell(row, 1).font = context.palette.font_small
    row += 2

    row = _section(ws, row, "CONCEPTOS", 6, context)
    headers = ["Codigo", "Descripcion", "Unidad", "Cantidad", "Precio", "Importe"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col)
        cell.value = header
        cell.fill = context.palette.fill_subhead
        cell.font = context.palette.font_subhead
        cell.alignment = Alignment(horizontal="center")
    row += 1

    line_by_code_desc = {
        (line.code, line.description): line
        for line in invoice.lines
    }
    for line_calc in calculation.line_calculations:
        line = line_by_code_desc[(line_calc.code, line_calc.description)]
        values = [
            line.code,
            line.description,
            line.unit,
            line.quantity,
            line.unit_price,
            line_calc.taxable_amount,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col)
            cell.value = value
            cell.border = BORDER
            if col in (5, 6):
                _format_money_cell(cell)
            elif col == 4:
                cell.number_format = QTY_FMT
        row += 1

    row += 1
    row = _section(ws, row, "RESUMEN FISCAL", 6, context)
    summary_rows = [
        ("Base imponible", calculation.taxable_base),
        (f"IVA ({calculation.vat.rate:.2%})", calculation.vat.amount),
    ]
    if calculation.vat.treatment in {TaxTreatment.EXEMPT, TaxTreatment.NOT_SUBJECT}:
        summary_rows[-1] = (f"IVA ({calculation.vat.treatment.value})", calculation.vat.amount)
    if calculation.withholding.amount:
        summary_rows.append((f"Retencion ({calculation.withholding.rate:.2%})", -calculation.withholding.amount))
    summary_rows.append(("TOTAL FACTURA", calculation.invoice_total))
    if calculation.withholding.amount:
        summary_rows.append(("A PAGAR", calculation.amount_due))

    for label, value in summary_rows:
        ws.cell(row, 4).value = label
        ws.cell(row, 4).font = Font(bold=label in {"TOTAL FACTURA", "A PAGAR"})
        ws.cell(row, 6).value = value
        _format_money_cell(ws.cell(row, 6))
        if label in {"TOTAL FACTURA", "A PAGAR"}:
            ws.cell(row, 4).fill = _fill(context.palette.color_fondo_alt)
            ws.cell(row, 6).fill = _fill(context.palette.color_fondo_alt)
            ws.cell(row, 6).font = Font(bold=True)
        row += 1

    row += 1
    row = _section(ws, row, "FORMA DE PAGO", 6, context)
    payment_rows = [
        ("Forma de pago", invoice.payment_terms.method.value),
        ("Vencimiento", invoice.payment_terms.due_date.isoformat() if invoice.payment_terms.due_date else ""),
        ("IBAN", invoice.payment_terms.iban),
        ("BIC", invoice.payment_terms.bic),
        ("Titular", invoice.payment_terms.account_holder),
        ("Referencia", invoice.payment_terms.payment_reference),
    ]
    for label, value in payment_rows:
        if _visible(value) != "-":
            _write_label_value(ws, row, label, _visible(value))
            row += 1

    if invoice.notes:
        row += 1
        row = _section(ws, row, "NOTAS", 6, context)
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=6)
        ws.cell(row, 1).value = invoice.notes
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

    for row_cells in ws.iter_rows():
        for cell in row_cells:
            if cell.value is not None:
                font = copy(cell.font)
                font.name = context.palette.font_family
                cell.font = font

    apply_document_footer(ws, metadata, context, 6, invoice=True)
    configure_excel_printing(ws, orientation="portrait", repeat_row=None, fit_to_height=1)
    wb.save(output_path)
    return output_path
