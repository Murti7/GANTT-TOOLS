"""QA estructural de paquetes XLSX/OOXML."""

from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
import tempfile
import zipfile
from xml.etree import ElementTree

from openpyxl import load_workbook


REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
WB_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@dataclass(frozen=True)
class ExcelIntegrityReport:
    path: Path
    zip_valid: bool
    xml_valid: bool
    relationships_valid: bool
    defined_names_valid: bool
    sheet_views_valid: bool
    roundtrip_valid: bool
    issues: list[str] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return (
            self.zip_valid
            and self.xml_valid
            and self.relationships_valid
            and self.defined_names_valid
            and self.sheet_views_valid
            and self.roundtrip_valid
            and not self.issues
        )


def validate_xlsx_integrity(path: Path) -> ExcelIntegrityReport:
    """Valida integridad OOXML basica sin depender de Microsoft Excel."""
    issues: list[str] = []
    zip_valid = xml_valid = relationships_valid = defined_names_valid = sheet_views_valid = roundtrip_valid = True

    try:
        with zipfile.ZipFile(path) as archive:
            bad_file = archive.testzip()
            if bad_file:
                zip_valid = False
                issues.append(f"ZIP corrupt member: {bad_file}")
            names = set(archive.namelist())
            xml_files = [name for name in names if name.endswith(".xml")]
            for name in xml_files:
                try:
                    ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as exc:
                    xml_valid = False
                    issues.append(f"Invalid XML {name}: {exc}")
            for rels_name in [name for name in names if name.endswith(".rels")]:
                rels_root = ElementTree.fromstring(archive.read(rels_name))
                source_dir = _rels_source_dir(rels_name)
                for rel in rels_root.findall(f"{REL_NS}Relationship"):
                    if rel.attrib.get("TargetMode") == "External":
                        continue
                    target = rel.attrib.get("Target", "")
                    resolved = _resolve_target(source_dir, target)
                    if resolved not in names:
                        relationships_valid = False
                        issues.append(f"Broken relationship {rels_name} -> {target} ({resolved})")
            if "xl/workbook.xml" in names:
                wb_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
                for defined_name in wb_root.findall(f".//{WB_NS}definedName"):
                    text = defined_name.text or ""
                    if "#REF!" in text or text.strip() in {"", "''!"}:
                        defined_names_valid = False
                        issues.append(f"Invalid definedName: {text!r}")
            for sheet_name in sorted(name for name in names if name.startswith("xl/worksheets/") and name.endswith(".xml")):
                try:
                    sheet_root = ElementTree.fromstring(archive.read(sheet_name))
                except ElementTree.ParseError:
                    continue
                for issue in _sheet_view_issues(sheet_root, sheet_name):
                    sheet_views_valid = False
                    issues.append(issue)
    except zipfile.BadZipFile as exc:
        zip_valid = xml_valid = relationships_valid = defined_names_valid = sheet_views_valid = False
        issues.append(f"Bad ZIP: {exc}")

    try:
        workbook = load_workbook(path)
        with tempfile.TemporaryDirectory() as tmpdir:
            roundtrip_path = Path(tmpdir) / "roundtrip.xlsx"
            workbook.save(roundtrip_path)
            workbook.close()
            loaded = load_workbook(roundtrip_path)
            loaded.close()
    except Exception as exc:
        roundtrip_valid = False
        issues.append(f"Roundtrip failed: {exc}")

    return ExcelIntegrityReport(
        path=path,
        zip_valid=zip_valid,
        xml_valid=xml_valid,
        relationships_valid=relationships_valid,
        defined_names_valid=defined_names_valid,
        sheet_views_valid=sheet_views_valid,
        roundtrip_valid=roundtrip_valid,
        issues=issues,
    )


def _sheet_view_issues(sheet_root: ElementTree.Element, sheet_name: str) -> list[str]:
    issues: list[str] = []
    for view_index, sheet_view in enumerate(sheet_root.findall(f".//{WB_NS}sheetView"), start=1):
        pane = sheet_view.find(f"{WB_NS}pane")
        selections = sheet_view.findall(f"{WB_NS}selection")
        if pane is None:
            for selection in selections:
                if "pane" in selection.attrib:
                    issues.append(
                        f"{sheet_name} sheetView[{view_index}] selection references pane "
                        f"{selection.attrib['pane']!r} but no pane exists"
                    )
            continue

        valid_panes = _valid_selection_panes(pane)
        active_pane = pane.attrib.get("activePane")
        if active_pane and active_pane not in valid_panes:
            issues.append(
                f"{sheet_name} sheetView[{view_index}] pane activePane {active_pane!r} "
                f"not in {sorted(valid_panes)}"
            )
        for selection in selections:
            selection_pane = selection.attrib.get("pane")
            if selection_pane and selection_pane not in valid_panes:
                issues.append(
                    f"{sheet_name} sheetView[{view_index}] selection pane {selection_pane!r} "
                    f"not in {sorted(valid_panes)}"
                )
    return issues


def _valid_selection_panes(pane: ElementTree.Element) -> set[str]:
    has_x_split = float(pane.attrib.get("xSplit", "0") or 0) > 0
    has_y_split = float(pane.attrib.get("ySplit", "0") or 0) > 0
    if has_x_split and has_y_split:
        return {"topRight", "bottomLeft", "bottomRight"}
    if has_x_split:
        return {"topRight"}
    if has_y_split:
        return {"bottomLeft"}
    return {"topLeft"}


def _rels_source_dir(rels_name: str) -> PurePosixPath:
    path = PurePosixPath(rels_name)
    if path.name == ".rels" and path.parent == PurePosixPath("_rels"):
        return PurePosixPath(".")
    if path.parent.name == "_rels":
        return path.parent.parent
    return path.parent


def _resolve_target(source_dir: PurePosixPath, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    parts: list[str] = []
    for part in (source_dir / target).parts:
        if part in {"", "."}:
            continue
        if part == "..":
            if parts:
                parts.pop()
        else:
            parts.append(part)
    return "/".join(parts)
