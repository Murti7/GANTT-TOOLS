from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl import Workbook
from PIL import Image

from gantt.bc3.models import BrandingConfig, LogoAssetType
from gantt.bc3.parser import parse_bc3
from gantt.planning.calculator import calcular_planificacion
from gantt.planning.models import cargar_planificacion_yaml
from gantt.reporting.documents import DocumentMetadata, DocumentType
from gantt.reporting.analysis_charts import AnalysisChartsReport
from gantt.reporting.excel_exporter import exportar_analisis
from gantt.reporting.network_diagram import generar_diagrama_red
from gantt.reporting.presentation import build_presentation_context
from gantt.reporting.rendering import render_budget_header
from gantt.reporting.styles import build_palette


BC3_MINIMO = r"""~C|OBRA##||Proyecto mini|100.00||0|
~D|OBRA##
|A#\\1
\|
~C|A#||Capitulo A|100.00||0|
~D|A#
|A01\\2
\|
~C|A01|ud|Partida de ejemplo|50.00||0|
~D|A01
|MO-OFI1\\3
\|
~C|MO-OFI1|h|Oficial 1a|20.00||1|
"""


PLANIFICACION_YAML = """
proyecto:
  nombre: Proyecto mini
  fecha_inicio: '2026-01-05'
  horas_dia: 8
  dias_semana: 5
  plazo_contractual_dias: 20

escenarios:
  base:
    MO-OFI1: 1

tareas:
  T1:
    nombre: Ejecutar partida
    tipo: tarea
    capitulos_bc3: ['A#']
    dependencias: []
  FIN:
    nombre: Fin
    tipo: hito
    capitulos_bc3: []
    duracion_dias_fija: 0
    dependencias: [T1]
    es_fin_plazo: true
"""


def _presupuesto(tmp_path: Path):
    bc3 = tmp_path / "mini.bc3"
    bc3.write_text(BC3_MINIMO, encoding="latin-1")
    return parse_bc3(bc3)


def _planificacion(tmp_path: Path, presupuesto):
    yaml_path = tmp_path / "planificacion.yaml"
    yaml_path.write_text(PLANIFICACION_YAML, encoding="utf-8")
    parametros, escenarios, tareas, bandas = cargar_planificacion_yaml(yaml_path)
    return calcular_planificacion(presupuesto, parametros, tareas, escenarios[0], bandas)


def _png(tmp_path: Path) -> Path:
    path = tmp_path / "logo.png"
    Image.new("RGBA", (200, 80), (10, 40, 35, 255)).save(path)
    return path


def _metadata() -> DocumentMetadata:
    return DocumentMetadata(
        document_type=DocumentType.CHAPTER_SUMMARY,
        title="Resumen por capitulos",
        project_name="Auditoria energetica integral",
        site_name="Viding Fitness Calvia",
        client_name="Cliente demo",
        reference="BAF-VID-2026-001",
        revision="00",
    )


def test_build_palette_valida_color_hexadecimal():
    with pytest.raises(ValueError, match="Color hexadecimal"):
        build_palette(BrandingConfig(color_primario="verde"))


def test_presentation_context_expone_tema_idioma_y_paleta():
    company = BrandingConfig(
        nombre="Demo Company",
        idioma="ca",
        color_primario="#123456",
        color_secundario="#654321",
        color_acento="#336699",
    )

    context = build_presentation_context(company)

    assert context.language == "ca"
    assert context.theme_id == "demo-company"
    assert context.palette.color_primario == "#123456"
    assert context.palette.chart.primary == "#123456"
    assert context.palette.diagram.task_border == "#654321"


def test_palette_semantic_foregrounds_contrastan_con_fondos_dinamicos():
    dark = build_palette(
        BrandingConfig(
            color_primario="#102030",
            color_secundario="#204020",
            color_texto="#111111",
        )
    )
    light = build_palette(
        BrandingConfig(
            color_primario="#F2F2F2",
            color_secundario="#E8F5E9",
            color_texto="#111111",
        )
    )

    assert dark.font_header.color.rgb.endswith("FFFFFF")
    assert dark.font_subhead.color.rgb.endswith("FFFFFF")
    assert light.font_header.color.rgb.endswith("000000")
    assert light.font_subhead.color.rgb.endswith("000000")


def test_budget_header_lockup_no_duplica_issuer_en_brand_area(tmp_path: Path):
    ws = Workbook().active
    company = BrandingConfig(
        nombre="BAFRAS Engineering S.L.",
        logo_path=_png(tmp_path),
        logo_asset_type=LogoAssetType.LOCKUP_HORIZONTAL,
    )

    render_budget_header(ws, _metadata(), build_presentation_context(company), 3)

    assert len(ws._images) == 1
    assert ws["A1"].value is None
    assert ws["A2"].value is None
    assert ws["A3"].value.startswith("BAF-VID-2026-001")


def test_budget_header_symbol_permite_identidad_textual(tmp_path: Path):
    ws = Workbook().active
    company = BrandingConfig(
        nombre="Empresa Simbolo",
        web="example.com",
        logo_path=_png(tmp_path),
        logo_asset_type=LogoAssetType.SYMBOL,
    )

    render_budget_header(ws, _metadata(), build_presentation_context(company), 3)

    assert len(ws._images) == 1
    assert ws["A1"].value == "Empresa Simbolo"
    assert ws["A2"].value == "example.com"


def test_budget_header_sin_logo_usa_marca_textual():
    ws = Workbook().active
    company = BrandingConfig(nombre="Empresa Sin Logo")

    render_budget_header(ws, _metadata(), build_presentation_context(company), 3)

    assert len(ws._images) == 0
    assert ws["A1"].value == "Empresa Sin Logo"
    assert ws["A2"].value == "Ingenieria - Energia - Sistemas"


def test_analysis_charts_generan_png_desde_modelo_sin_excel(tmp_path: Path):
    presupuesto = _presupuesto(tmp_path)
    palette = build_palette(BrandingConfig(color_primario="#123456", color_secundario="#2A7A4F"))
    output_dir = tmp_path / "reporting"

    AnalysisChartsReport.from_presupuesto(presupuesto, output_dir, palette=palette).generate()

    assert (output_dir / "01_economico" / "01_importe_por_capitulo.png").stat().st_size > 0
    assert (output_dir / "02_mano_obra" / "01_horas_mo_por_capitulo.png").stat().st_size > 0


def test_exportar_analisis_aplica_paleta_en_cabecera(tmp_path: Path):
    presupuesto = _presupuesto(tmp_path)
    palette = build_palette(BrandingConfig(color_primario="#123456", color_secundario="#2A7A4F"))
    output_path = tmp_path / "analisis.xlsx"

    exportar_analisis(presupuesto, output_path, palette=palette)

    wb = load_workbook(output_path, read_only=False)
    try:
        assert wb["Resumen por capítulo"]["A1"].fill.fgColor.rgb.endswith("123456")
    finally:
        wb.close()


def test_diagrama_red_acepta_paleta_dinamica(tmp_path: Path):
    presupuesto = _presupuesto(tmp_path)
    planificacion = _planificacion(tmp_path, presupuesto)
    palette = build_palette(BrandingConfig(color_primario="#123456", color_secundario="#2A7A4F"))

    png = generar_diagrama_red(planificacion, palette=palette)

    assert png.startswith(b"\x89PNG")
    assert len(png) > 1000
