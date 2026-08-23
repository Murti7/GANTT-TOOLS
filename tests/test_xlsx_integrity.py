from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl import Workbook

from gantt.application.budgeting import run_budgeting
from gantt.application.context import resolver_contexto_ejecucion
from gantt.billing.calculator import calculate_invoice
from gantt.billing.models import Client, InvoiceLine, Issuer
from gantt.billing.services import create_invoice
from gantt.reporting.invoice_exporter import export_invoice_excel
from gantt.reporting.presentation import build_presentation_context
from gantt.reporting.rendering import configure_excel_printing
from gantt.reporting.xlsx_integrity import validate_xlsx_integrity


BC3_MINIMO = r"""~C|OBRA##||Proyecto mini|100.00||0|
~D|OBRA##
|A#\\1
\|
~C|A#||Capitulo A|100.00||0|
~D|A#
|A01\\2
\|
~C|A01|ud|Partida de ejemplo|50.00||0|
~D|A01
|MO-OFI1\\3
\|
~C|MO-OFI1|h|Oficial 1a|20.00||1|
"""


def test_excel_integrity_gate_valida_invoice_roundtrip(tmp_path: Path):
    invoice = create_invoice(
        Issuer(legal_name="BAFRAS Engineering S.L.", tax_id="B123", fiscal_address="Calle Demo"),
        Client(legal_name="Cliente S.A.", tax_id="A123", fiscal_address="Avenida Cliente"),
        [InvoiceLine(description="Servicio", unit_price=Decimal("1000"))],
    )
    path = export_invoice_excel(
        invoice,
        calculate_invoice(invoice),
        build_presentation_context(),
        tmp_path / "invoice.xlsx",
    )

    report = validate_xlsx_integrity(path)

    assert report.passed, report.issues


def test_excel_integrity_gate_detecta_selection_pane_huerfano(tmp_path: Path):
    workbook = Workbook()
    ws = workbook.active
    ws["A1"] = "demo"
    ws.freeze_panes = "A8"
    ws.freeze_panes = None
    path = tmp_path / "orphan-pane.xlsx"
    workbook.save(path)

    report = validate_xlsx_integrity(path)

    assert not report.sheet_views_valid
    assert any("selection references pane" in issue for issue in report.issues)


def test_delivery_printing_no_crea_selection_pane_huerfano(tmp_path: Path):
    workbook = Workbook()
    ws = workbook.active
    ws["A1"] = "demo"

    configure_excel_printing(ws, repeat_row=1)
    path = tmp_path / "delivery.xlsx"
    workbook.save(path)

    report = validate_xlsx_integrity(path)

    assert report.passed, report.issues


def test_budget_analysis_autofilter_no_cubre_cabeceras_fusionadas(tmp_path: Path):
    input_dir = tmp_path / "projects" / "Mini" / "input"
    input_dir.mkdir(parents=True)
    (input_dir / "mini.bc3").write_text(BC3_MINIMO, encoding="latin-1")
    (input_dir / "config.yaml").write_text("proyecto: Proyecto mini\n", encoding="utf-8")

    result = run_budgeting(resolver_contexto_ejecucion("Mini", workspace_root=tmp_path))
    analysis = result.output_root / "analysis" / "budget_analysis.xlsx"

    report = validate_xlsx_integrity(analysis)
    assert report.passed, report.issues

    wb = load_workbook(analysis)
    try:
        ws = wb.active
        assert ws.auto_filter.ref.startswith("A")
        merged_rows = {row for merged in ws.merged_cells.ranges for row, _col in merged.cells}
        filter_start_row = int(ws.auto_filter.ref.split(":")[0][1:])
        assert filter_start_row not in merged_rows
    finally:
        wb.close()
